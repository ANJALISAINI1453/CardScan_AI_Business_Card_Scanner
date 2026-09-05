from flask import Flask, request, jsonify, send_file, send_from_directory
from dotenv import load_dotenv
from PIL import Image, ImageOps, ImageEnhance, ImageFilter
from google import genai
from google.genai import types
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
import os, io, json
import gspread
from google.oauth2.service_account import Credentials

load_dotenv()
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)
EXCEL_FILE = DATA_DIR / "Business_Cards.xlsx"
API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
MODEL = os.getenv("GEMINI_MODEL", "gemini-3.6-flash").strip()
GOOGLE_SHEET_NAME = os.getenv("GOOGLE_SHEET_NAME", "Business Cards Database").strip()
GOOGLE_WORKSHEET_NAME = os.getenv("GOOGLE_WORKSHEET_NAME", "Business Cards").strip()
GOOGLE_SHEET_URL = os.getenv("GOOGLE_SHEET_URL", "").strip()
SERVICE_ACCOUNT_FILE = BASE_DIR / os.getenv("GOOGLE_SERVICE_ACCOUNT_FILE", "credentials/service_account.json").strip()
GOOGLE_SERVICE_ACCOUNT_JSON = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()
app = Flask(__name__)

class BusinessCardData(BaseModel):
    name: str = ""
    business_name: str = ""
    designation: str = ""
    phone: str = ""
    alternate_phone: str = ""
    email: str = ""
    website: str = ""
    street_address: str = ""
    locality: str = ""
    district: str = ""
    state: str = ""
    pin: str = ""
    owner_details: str = ""
    other_information: str = ""

HEADERS = ["Scan Date & Time","Name","Business Name","Designation","Phone","Alternate Phone","Email","Website","Street Address","Area / Locality","District","State","PIN / ZIP","Owner / Partner Details","Other Information"]

def now_string():
    return datetime.now(ZoneInfo("Asia/Kolkata")).strftime("%d-%m-%Y %I:%M:%S %p")

GOOGLE_SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]

def google_sheets_configured():
    return bool(GOOGLE_SERVICE_ACCOUNT_JSON) or SERVICE_ACCOUNT_FILE.exists()

def get_google_client():
    if GOOGLE_SERVICE_ACCOUNT_JSON:
        try:
            service_account_info = json.loads(GOOGLE_SERVICE_ACCOUNT_JSON)
        except json.JSONDecodeError:
            raise RuntimeError("GOOGLE_SERVICE_ACCOUNT_JSON contains invalid JSON.")

        creds = Credentials.from_service_account_info(
            service_account_info,
            scopes=GOOGLE_SCOPES
        )
    else:
        if not SERVICE_ACCOUNT_FILE.exists():
            raise RuntimeError(
                f"Google service-account file not found: {SERVICE_ACCOUNT_FILE}. "
                "Put your downloaded service_account.json inside the credentials folder."
            )

        creds = Credentials.from_service_account_file(
            str(SERVICE_ACCOUNT_FILE),
            scopes=GOOGLE_SCOPES
        )

    return gspread.authorize(creds)

def get_google_worksheet():
    gc = get_google_client()
    if GOOGLE_SHEET_URL:
        sh = gc.open_by_url(GOOGLE_SHEET_URL)
    else:
        try:
            sh = gc.open(GOOGLE_SHEET_NAME)
        except gspread.SpreadsheetNotFound:
            sh = gc.create(GOOGLE_SHEET_NAME)
    try:
        ws = sh.worksheet(GOOGLE_WORKSHEET_NAME)
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title=GOOGLE_WORKSHEET_NAME, rows=1000, cols=len(HEADERS))
    # Make sure the first row contains the required headers. Existing data is preserved.
    first_row = ws.row_values(1)
    if first_row != HEADERS:
        if not first_row:
            ws.update("A1", [HEADERS])
        else:
            # Only set headers if the first row is not already our headers and the sheet is effectively empty.
            values = ws.get_all_values()
            if len(values) <= 1 and not any(first_row):
                ws.update("A1", [HEADERS])
    return ws

def save_to_google_sheet(data):

    ws = get_google_worksheet()

    row = [
        data.get("scan_datetime") or now_string(),
        data.get("name", ""),
        data.get("business_name", ""),
        data.get("designation", ""),
        data.get("phone", ""),
        data.get("alternate_phone", ""),
        data.get("email", ""),
        data.get("website", ""),
        data.get("street_address", ""),
        data.get("locality", ""),
        data.get("district", ""),
        data.get("state", ""),
        data.get("pin", ""),
        data.get("owner_details", ""),
        data.get("other_information", "")
    ]

    # Find the next empty row based only on Column A
    col_a = ws.col_values(1)

    next_row = len(col_a) + 1

    # Write the complete record into A:O of the new row
    ws.update(
        f"A{next_row}:O{next_row}",
        [row],
        value_input_option="USER_ENTERED"
    )

    return ws.spreadsheet.url

def ensure_excel():
    if EXCEL_FILE.exists(): return
    wb = Workbook()
    ws = wb.active
    ws.title = "Business Cards"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center")
    widths = [22,25,30,25,18,20,35,35,45,25,25,25,15,35,45]
    for i,w in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(EXCEL_FILE)

def preprocess(img):
    img = ImageOps.exif_transpose(img)
    if img.mode != "RGB": img = img.convert("RGB")
    if img.width < 1800:
        ratio = 1800 / img.width
        img = img.resize((1800, int(img.height*ratio)), Image.Resampling.LANCZOS)
    img = ImageEnhance.Contrast(img).enhance(1.18)
    img = ImageEnhance.Sharpness(img).enhance(1.35)
    return img.filter(ImageFilter.SHARPEN)

def extract_card(img):
    if not API_KEY:
        raise RuntimeError("GEMINI_API_KEY is missing in .env")
    client = genai.Client(api_key=API_KEY)
    img = preprocess(img)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=95)
    part = types.Part.from_bytes(data=buf.getvalue(), mime_type="image/jpeg")
    prompt = """Read this business card carefully and extract all visible business/contact information.
Return ONLY the requested JSON schema. Never guess or invent.
Name = person's full name exactly as printed.
Business name = company/firm/shop/organization.
Designation = role/position/qualification when applicable.
Phone and alternate_phone = phone/mobile numbers only, preserving digits exactly.
Email and website = exact visible values.
Split address into street_address, locality, district, state and pin only when supported by the card.
owner_details = ONLY explicit owner/partner/proprietor information; never infer it.
other_information = useful remaining printed information such as qualifications, branch/head office, GST/registration/fax.
If absent, use an empty string.
Never confuse PIN, GST or registration numbers with phone numbers."""
    response = client.models.generate_content(
        model=MODEL, contents=[part, prompt],
        config=types.GenerateContentConfig(
            temperature=0, response_mime_type="application/json",
            response_schema=BusinessCardData
        )
    )
    if getattr(response, "parsed", None):
        return response.parsed.model_dump()
    if getattr(response, "text", None):
        return json.loads(response.text)
    raise RuntimeError("AI returned an empty response.")

@app.route("/")
def index():
    return send_from_directory(BASE_DIR, "index.html")

@app.route("/scan", methods=["POST"])
def scan():
    try:
        f = request.files.get("image")
        if not f or not f.filename:
            return jsonify(success=False, error="Please select a card image."), 400
        img = Image.open(f.stream)
        data = extract_card(img)
        return jsonify(success=True, data=data, scan_datetime=now_string())
    except Exception as e:
        print("SCAN ERROR:", repr(e))
        return jsonify(success=False, error=str(e)), 500

@app.route("/save", methods=["POST"])
def save():
    try:
        data = request.get_json(silent=True) or {}
        ensure_excel()
        wb = load_workbook(EXCEL_FILE)
        ws = wb["Business Cards"]
        ws.append([
            data.get("scan_datetime") or now_string(),
            data.get("name",""), data.get("business_name",""), data.get("designation",""),
            data.get("phone",""), data.get("alternate_phone",""), data.get("email",""),
            data.get("website",""), data.get("street_address",""), data.get("locality",""),
            data.get("district",""), data.get("state",""), data.get("pin",""),
            data.get("owner_details",""), data.get("other_information","")
        ])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        wb.save(EXCEL_FILE)
        return jsonify(success=True, message="Saved to Excel successfully.")
    except Exception as e:
        print("SAVE ERROR:", repr(e))
        return jsonify(success=False, error=str(e)), 500

@app.route("/save-google", methods=["POST"])
def save_google():
    try:
        data = request.get_json(silent=True) or {}
        url = save_to_google_sheet(data)
        return jsonify(success=True, message="Saved to Google Sheet successfully.", sheet_url=url)
    except Exception as e:
        print("GOOGLE SHEETS ERROR:", repr(e))
        return jsonify(success=False, error=str(e)), 500

@app.route("/google-status")
def google_status():
    return jsonify(
        configured=google_sheets_configured(),
        sheet_name=GOOGLE_SHEET_NAME,
        worksheet_name=GOOGLE_WORKSHEET_NAME,
        sheet_url=GOOGLE_SHEET_URL,
        service_account_file=str(SERVICE_ACCOUNT_FILE),
    )

@app.route("/download")
def download():
    ensure_excel()
    return send_file(EXCEL_FILE, as_attachment=True, download_name="Business_Cards.xlsx")

@app.route("/health")
def health():
    return jsonify(ok=True, api_key_configured=bool(API_KEY), model=MODEL, google_sheets_configured=google_sheets_configured())

if __name__ == "__main__":
    ensure_excel()
    print("\nCardScan AI running at http://127.0.0.1:5000\n")
    app.run(host="127.0.0.1", port=5000, debug=True)